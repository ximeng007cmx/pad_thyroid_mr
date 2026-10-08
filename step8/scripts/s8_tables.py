# -*- coding: utf-8 -*-
"""
Step 8 共定位结果 → XLSX 数据表。
数值全部来自 step8/results/*.tsv / *.json，不手抄。
输出：deliverables/Step8_共定位结果_数据表.xlsx
"""
import io
import os
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = r"D:\WorkSpace_Study\pad_thyroid_mr"
RES = os.path.join(ROOT, "step8", "results")
DELIV = os.path.join(ROOT, "deliverables")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def load_tsv(p):
    with io.open(p, encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autowidth(ws):
    for col in ws.columns:
        w = 10
        for cell in col:
            v = cell.value
            if v is None:
                continue
            ln = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
            w = max(w, min(ln + 2, 46))
        ws.column_dimensions[col[0].column_letter].width = w


rows = load_tsv(os.path.join(RES, "coloc_results.tsv"))
cov = load_tsv(os.path.join(RES, "precheck_coverage.tsv"))
with io.open(os.path.join(RES, "locus_annotation.json"), encoding="utf-8") as f:
    ANN = json.load(f)


def gene(rid):
    g = [x for x in (ANN.get(rid, {}).get("genes") or []) if not x.startswith("ENSG")]
    return "/".join(g) if g else "—"


def fl(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


wb = Workbook()

# ---------------- 表1：共定位结果
ws = wb.active
ws.title = "共定位结果"
cols = ["位点", "邻近基因", "区域(GRCh38)", "rsID", "MHC", "结局", "可配对变异",
        "匹配(rsID)", "匹配(位置)", "PP.H0", "PP.H1", "PP.H2", "PP.H3", "PP.H4", "判定",
        "H4(sd=0.20)", "H4(p12=1e-6)", "H4(p12=1e-4)"]
ws.append(cols)
for r in rows:
    ws.append([r["region"], gene(r["region"]),
               "chr%s:%s" % (r["chr"], r["pos"]), r.get("rs", ""),
               "是" if str(r.get("mhc")).lower() in ("1", "true") else "",
               r["outcome_label"], int(r.get("n") or 0),
               int(r.get("m_rs") or 0), int(r.get("m_pos") or 0),
               fl(r.get("H0")), fl(r.get("H1")), fl(r.get("H2")), fl(r.get("H3")),
               fl(r.get("H4")), r.get("verdict", ""),
               fl(r.get("H4_sd020")), fl(r.get("H4_p12_1e6")), fl(r.get("H4_p12_1e4"))])
style_header(ws, len(cols))
autowidth(ws)

# ---------------- 表2：前置覆盖度
ws2 = wb.create_sheet("前置覆盖度")
TAG_MAP = {"thy_hypo": "甲减（暴露）", "thy_tsh": "TSH", "thy_ft4": "FT4",
           "cad_nouk": "CAD-Nikpay", "is_sv": "IS-SV", "is_laa": "IS-LAA",
           "mvp_meta": "MVP-PAD"}
cols2 = ["数据集", "每区变异(中位)", "min", "max", "lead 命中", "坐标口径"]
ws2.append(cols2)
for k in cov[0].keys():
    if not k.startswith("n_"):
        continue
    tag = k[2:]
    vals = sorted(int(r[k]) for r in cov)
    ws2.append([TAG_MAP.get(tag, tag), vals[len(vals) // 2], vals[0], vals[-1],
                "%d/%d" % (sum(1 for r in cov if str(r.get("lead_" + tag)).lower() in ("true", "1")), len(cov)),
                "GRCh37(hg19) → 仅 rsID 匹配" if "mvp" in tag else "GRCh38"])
style_header(ws2, len(cols2))
autowidth(ws2)

# ---------------- 表3：位点级配对数量
ws3 = wb.create_sheet("位点-结局配对")
cols3 = ["位点", "邻近基因", "区域", "结局", "甲减区变异", "结局区变异", "可配对"]
ws3.append(cols3)
for r in rows:
    ws3.append([r["region"], gene(r["region"]), "chr%s:%s" % (r["chr"], r["pos"]),
                r["outcome_label"], "", int(r.get("n_out_region") or 0), int(r.get("n") or 0)])
style_header(ws3, len(cols3))
autowidth(ws3)

out = os.path.join(DELIV, "Step8_共定位结果_数据表.xlsx")
wb.save(out)
print("OK  " + out, os.path.getsize(out), "B")
