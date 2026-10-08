# -*- coding: utf-8 -*-
"""
Step 9 统计功效结果 → XLSX 数据表。
数值全部来自 step9/results/*.tsv / *.json，不手抄。
输出：deliverables/Step9_统计功效_数据表.xlsx
"""
import io
import os
import csv
import json

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT = r"D:\WorkSpace_Study\pad_thyroid_mr"
RES = os.path.join(ROOT, "step9", "results")
DELIV = os.path.join(ROOT, "deliverables")

HDR_FILL = PatternFill("solid", fgColor="1F4E79")
HDR_FONT = Font(bold=True, color="FFFFFF", size=10)
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BAD_FILL = PatternFill("solid", fgColor="FDEBEB")

TRAIT_CN = {"hypo": "甲减", "tsh": "TSH", "ft4": "FT4"}
TRAIT_ORDER = ["hypo", "tsh", "ft4"]
NOTE_RANK = {"无重叠": 0, "主分析·含UKB": 1, "重叠": 2, "重叠·甲减失效": 3}
KEY_CN = {"sakaue": "PAD（Sakaue 2021）", "finngen": "PAD（FinnGen R13）",
          "mvp_meta": "PAD（MVP trans-ancestry）", "cad_nouk": "CAD（Nikpay 2015）",
          "cad_ukb": "CAD（UKB Mbatchou 2021）", "mi_ukb": "MI（UKB Donertas 2021）",
          "is_laa": "IS-LAA（MEGASTROKE）", "is_sv": "IS-SV（MEGASTROKE）"}


def load_tsv(p):
    with io.open(p, encoding="utf-8") as f:
        return list(csv.DictReader(f, delimiter="\t"))


def short(title):
    if "（" not in title:
        return title
    name, rest = title.split("（", 1)
    return "%s（%s）" % (name, rest.rstrip("）").split(",")[0].strip())


def fl(x):
    if x is None:
        return None
    s = str(x).strip()
    if s == "":
        return None
    try:
        return float(s)
    except ValueError:
        return None


def style_header(ws, ncol, row=1):
    for c in range(1, ncol + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autowidth(ws):
    for col in ws.columns:
        w = 10
        for cell in col:
            v = cell.value
            if v is None:
                continue
            ln = sum(2 if ord(ch) > 127 else 1 for ch in str(v))
            w = max(w, min(ln + 2, 46))
        ws.column_dimensions[col[0].column_letter].width = w


FW = load_tsv(os.path.join(RES, "power_forward.tsv"))
RV = load_tsv(os.path.join(RES, "power_reverse.tsv"))
MV = load_tsv(os.path.join(RES, "power_mvmr.tsv"))
SCALE = load_tsv(os.path.join(RES, "scale_check.tsv"))
with io.open(os.path.join(RES, "power_summary.json"), encoding="utf-8") as f:
    SUM = json.load(f)
INST, SCEN = SUM["instruments"], SUM["scenarios"]

SK = lambda r: (TRAIT_ORDER.index(r["exposure_key"]),
                NOTE_RANK.get(r["source_note"], 9), r["outcome"])
FW, MV = sorted(FW, key=SK), sorted(MV, key=SK)

wb = Workbook()

# ---------------- 表1：工具强度
ws = wb.active
ws.title = "工具强度"
cols = ["暴露", "工具数", "累计 R²", "中位 F", "最小 F", "有效样本量"]
ws.append(cols)
for t in TRAIT_ORDER:
    s = INST[t]
    ws.append([TRAIT_CN[t], s["n"], round(s["r2"], 6), round(s["f_median"], 2),
               round(s["f_min"], 2), int(round(s["n_eff"]))])
style_header(ws, len(cols))
autowidth(ws)

# ---------------- 表2：结局尺度体检
ws2 = wb.create_sheet("结局尺度体检")
cols2 = ["结局文件", "仪器 SNP 数", "中位 SE", "参考中位 SE", "比值", "判定"]
ws2.append(cols2)
for s in SCALE:
    bad = bool(str(s.get("flag") or "").strip())
    ws2.append([KEY_CN.get(s["key"], s["key"]), int(s["n_se"]), fl(s["med_se"]),
                fl(s["ref_med_se"]), round(fl(s["ratio_to_ref"]), 4),
                s["flag"] or "OK · log-OR 尺度"])
    if bad:
        for c in range(1, len(cols2) + 1):
            ws2.cell(row=ws2.max_row, column=c).fill = BAD_FILL
style_header(ws2, len(cols2))
autowidth(ws2)

# ---------------- 表3：正向 MR 功效
ws3 = wb.create_sheet("正向MR-功效")
cols3 = ["暴露", "结局", "重叠", "nSNP", "观测 β", "p(固定)", "MDE80(OR)", "MDE90(OR)",
         "功效@OR1.05", "功效@OR1.10", "功效@OR1.15", "尺度", "单位"]
ws3.append(cols3)
for r in FW:
    bad = bool(str(r.get("scale_flag") or "").strip())
    ws3.append([r["exposure"], short(r["outcome"]), r["source_note"], int(r["nsnp"]),
                fl(r["beta"]), fl(r["p_fixed"]),
                None if bad else fl(r["mde80_or"]), None if bad else fl(r["mde90_or"]),
                None if bad else fl(r["pow_or1.05"]), None if bad else fl(r["pow_or1.10"]),
                None if bad else fl(r["pow_or1.15"]),
                r["scale_flag"] or "log-OR", r["unit"]])
    if bad:
        for c in range(1, len(cols3) + 1):
            ws3.cell(row=ws3.max_row, column=c).fill = BAD_FILL
style_header(ws3, len(cols3))
autowidth(ws3)

# ---------------- 表4：反向 MR 功效
ws4 = wb.create_sheet("反向MR-功效")
cols4 = ["甲状腺结局", "血管暴露", "重叠", "nSNP", "观测 β", "p(固定)",
         "MDE80(OR)", "MDE80(SD)", "结局类型", "单位"]
ws4.append(cols4)
for r in sorted(RV, key=lambda x: (TRAIT_ORDER.index(x["exposure_key"]), x["outcome"])):
    ws4.append([r["exposure"], short(r["outcome"]), r["source_note"], int(r["nsnp"]),
                fl(r["beta"]), fl(r["p_fixed"]), fl(r["mde80_or"]), fl(r["mde80_sd"]),
                r["outcome_type"], r["unit"]])
style_header(ws4, len(cols4))
autowidth(ws4)

# ---------------- 表5：MVMR 功效
ws5 = wb.create_sheet("MVMR-功效")
cols5 = ["结局", "重叠", "nSNP", "条件 F", "观测 β", "p", "MDE80(OR)", "尺度", "单位"]
ws5.append(cols5)
for r in MV:
    if r["exposure_key"] != "hypo":
        continue
    bad = bool(str(r.get("scale_flag") or "").strip())
    ws5.append([short(r["outcome"]), r["source_note"], int(r["nsnp"]), fl(r["condF"]),
                fl(r["beta"]), fl(r["p_fixed"]), None if bad else fl(r["mde80_or"]),
                r["scale_flag"] or "log-OR", r["unit"]])
    if bad:
        for c in range(1, len(cols5) + 1):
            ws5.cell(row=ws5.max_row, column=c).fill = BAD_FILL
style_header(ws5, len(cols5))
autowidth(ws5)

# ---------------- 表6：情景分析
ws6 = wb.create_sheet("情景-样本量放大")
cols6 = ["暴露", "结局（主分析 PAD, Sakaue 2021）", "×1", "×2", "×5", "×10", "说明"]
ws6.append(cols6)
for t in TRAIT_ORDER:
    k = "%s|PAD（Sakaue 2021, GCST90018890）" % t
    ws6.append([TRAIT_CN[t], "MDE80 (OR)", fl(SCEN["×1"].get(k)), fl(SCEN["×2"].get(k)),
                fl(SCEN["×5"].get(k)), fl(SCEN["×10"].get(k)), "SE ∝ 1/√n"])
style_header(ws6, len(cols6))
autowidth(ws6)

# ---------------- 表7：说明
ws7 = wb.create_sheet("口径与说明")
notes = [
    ["项目", "说明"],
    ["功效公式", "Power(β_true)=Φ(λ−z_{1−α/2})+Φ(−λ−z_{1−α/2})，λ=|β_true|/SE；"
                "MDE(power)=(z_{1−α/2}+z_power)×SE"],
    ["显著性水平", "α=0.05（双侧），z=1.95996"],
    ["功效档", "80% → z=0.84162；90% → z=1.28155"],
    ["SE 来源", "直接取自已完成的 MR 结果文件（IVW 固定 / 随机），不重新估计效应"],
    ["单位（正向 / MVMR）", "甲减＝每「甲减 odds 翻倍」（ln2）；TSH / FT4＝每 1 SD；结局二分 → 报 OR"],
    ["单位（反向）", "暴露＝血管疾病（二分）→「每血管疾病 odds 翻倍」；"
                    "甲减结局报 OR，TSH / FT4 结局报 SD 变化"],
    ["MDE 含义", "在给定功效下可检出的最小效应；观测效应若 ≥ MDE 则本就在检出能力内"],
    ["不使用 post-hoc power", "负研究「观测效应的功效」恒偏低，会系统性把阴性误判为功效不足，故不作判据"],
    ["结局尺度体检", "以仪器 SNP 的中位 se_Y 横向比对，偏离参考中位 ≥5 倍判为尺度异常；"
                    "常数尺度重标定不改变 z 与 p，仅效应量受影响"],
    ["MI（UKB, GCST90038610）", "判定为线性概率尺度（中位 SE 为参考的 0.030 倍）；"
                              "本次不改动原始数据，仅将其 MDE 置空、不报 OR"],
    ["情景分析假设", "SE ∝ 1/√n，其余不变；未考虑病例比例与芯片覆盖变化"],
    ["脚本", "step9/scripts/s9_power.py（计算）、step9/scripts/s9_tables.py（本表）、"
             "step9/scripts/s9_onepager.py（一页纸）"],
    ["生成日期", "2026-10-07"],
]
for row in notes:
    ws7.append(row)
style_header(ws7, 2)
autowidth(ws7)

out = os.path.join(DELIV, "Step9_统计功效_数据表.xlsx")
wb.save(out)
print("OK  " + out, os.path.getsize(out), "B")
